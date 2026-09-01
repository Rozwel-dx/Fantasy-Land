import argparse
from ast import main
import logging
import os
import re
from collections import defaultdict, namedtuple
from io import BytesIO
from typing import Dict, List, Tuple, Any

import matplotlib.pyplot as plt
import pandas as pd
from decimal import Decimal, getcontext
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter

# 设置matplotlib不使用GUI后端
plt.switch_backend('Agg')

# excel表格最大列宽
MAX_COL_WIDTH = 20
# task_summary_by_process_pid表每张图中最多展示20个pid，其余部分进行拆分
MAX_PID_PER_CHART = 20
# task_summary_by_process_pid表每种comm图的起始列
BASE_COL_INDEX = 12
ROW_HEIGHT = 15

TaskKey = namedtuple("TaskKey", ["cpu", "comm", "pid"])

CPU_RE = re.compile(r'\[(\d+)]')
TS_RE = re.compile(r'\s(\d+\.\d+):')
IRQ_ENTRY_RE = re.compile(r'irq_handler_entry: irq=(\d+)\s+name=(\S+)')
IRQ_EXIT_RE = re.compile(r'irq_handler_exit: irq=(\d+)\s+ret=')
SOFT_ENTRY_RE = re.compile(r'softirq_entry: vec=(\d+)\s+\[action=([^]]+)]')
SOFT_EXIT_RE = re.compile(r'softirq_exit: vec=(\d+)\s+\[action=([^]]+)]')
WAKING_RE = re.compile(r'sched_waking: comm=([^ ]+) pid=(\d+) prio=\d+ target_cpu=(\d+)')
WAKEUP_RE = re.compile(r'sched_wakeup: comm=([^ ]+) pid=(\d+) prio=\d+ success=\d+ target_cpu=(\d+)')
SWITCH_RE = re.compile(
    r'sched_switch: prev_comm=([^ ]+) prev_pid=(\d+) prev_prio=\d+ prev_state=([A-Z0-9]+) '
    r'==> next_comm=([^ ]+) next_pid=(\d+) next_prio=\d+'
)
MIGRATE_RE = re.compile(
    r'sched_migrate_task: comm=([^ ]+) pid=(\d+) prio=\d+ orig_cpu=(\d+) dest_cpu=(\d+)'
)


class TaskStats:
    __slots__ = ("running", "sleeping", "runnable", "cs_count", "preempt_cs_count", "preempt_time_s", "migrate_count", "irqs")

    def __init__(self):
        self.running = Decimal("0")
        self.sleeping = Decimal("0")
        self.runnable = Decimal("0")
        self.cs_count = 0
        self.preempt_cs_count = 0
        self.preempt_time_s = Decimal("0")
        self.migrate_count = 0
        self.irqs = defaultdict(lambda: {"count": 0, "time": Decimal("0")})


class OffCPUState:
    __slots__ = ("state", "timestamp", "cpu")

    def __init__(self, state, timestamp, cpu):
        self.state = state
        self.timestamp = timestamp
        self.cpu = cpu


class IRQFrame:
    __slots__ = ("irq_type", "irq_id", "name", "start_ts", "task_key", "acc_time")

    def __init__(self, irq_type, irq_id, name, start_ts, task_key):
        self.irq_type = irq_type
        self.irq_id = irq_id
        self.name = name
        self.start_ts = start_ts
        self.task_key = task_key
        self.acc_time = Decimal("0")


class TraceAnalyzer:
    def __init__(self):
        self.current_task = {}
        self.last_ts = defaultdict(lambda: None)
        self.irq_stack = defaultdict(list)
        self.stats = defaultdict(TaskStats)
        self.offcpu = {}
        self.pid_comm = {}

    @staticmethod
    def classify_prev_state(prev_state_str: str) -> str:
        if any(ch in prev_state_str for ch in ("S", "D", "K", "W", "T")):
            return "sleeping"
        elif "R" in prev_state_str:
            return "preempted"
        else:
            return "idle"

    @staticmethod
    def get_task_key(cpu: int, comm: str, pid: int) -> TaskKey:
        return TaskKey(cpu=cpu, comm=comm, pid=int(pid))

    @staticmethod
    def write_excel(
            df_summary: pd.DataFrame,
            df_irq_detail: pd.DataFrame,
            df_proc_irq: pd.DataFrame,
            df_summary_by_comm: pd.DataFrame,
            df_summary_by_process_pid: pd.DataFrame,
            path: str
    ) -> None:
        # 定义工作表名称和对应的数据帧
        sheets = {
            "task_summary": df_summary,
            "task_summary_by_comm": df_summary_by_comm,
            "task_summary_by_process_pid": df_summary_by_process_pid,
            "task_irq_detail": df_irq_detail,
            "proc_irq_detail": df_proc_irq
        }
        
        # 写入Excel文件
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for sheet_name, df in sheets.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        # 调整列宽并添加筛选器
        wb = load_workbook(path)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws.auto_filter.ref = ws.dimensions
            
            # 计算并设置列宽
            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    cell_value = str(cell.value) if cell.value is not None else ""
                    if len(cell_value) > max_length:
                        max_length = len(cell_value)
                ws.column_dimensions[col_letter].width = min(max_length + 2, MAX_COL_WIDTH)
        
        wb.save(path)


    @staticmethod
    def create_bar_chart(df: pd.DataFrame, x_col: str, y_col: str, title: str, 
                        y_label: str, color: str = "steelblue", 
                        font_size: int = 8, dpi: int = 100) -> BytesIO:
        """创建柱状图并返回BytesIO对象"""
        plt.figure(figsize=(6, 4))
        ax = plt.gca()
        ax.bar(df[x_col], df[y_col], color=color)
        ax.set_ylabel(y_label)
        ax.set_title(title)
        ax.tick_params(axis="x", rotation=45)
        ax.margins(y=0.1)
        
        # 添加数值标签
        for x, y in zip(df[x_col], df[y_col]):
            ax.text(x, y, f"{int(y)}", ha="center", va="bottom", fontsize=font_size)
        
        plt.tight_layout()
        buf = BytesIO()
        plt.savefig(buf, format="png", dpi=dpi)
        plt.close()
        buf.seek(0)
        return buf
    
    @staticmethod
    def create_subplots_chart(df_list: List[pd.DataFrame], x_col: str, y_cols: List[str], 
                            titles: List[str], y_labels: List[str], 
                            main_title: str, nrows: int, ncols: int, 
                            figsize: Tuple[int, int] = (8, 6), color: Any = "royalblue",
                            font_size: int = 6, dpi: int = 100) -> BytesIO:
        """创建包含多个子图的图表并返回BytesIO对象"""
        fig, axes = plt.subplots(nrows, ncols, figsize=figsize)
        fig.suptitle(main_title, fontsize=12)
        
        # 确保axes是可迭代的
        if nrows == 1 and ncols == 1:
            axes = [axes]
        else:
            axes = axes.flatten()
        
        # 处理颜色参数，可以是单个颜色字符串或颜色列表
        colors = [color] * len(df_list) if isinstance(color, str) else color
        
        for i, (ax, df, y_col, title, y_label, bar_color) in enumerate(zip(axes, df_list, y_cols, titles, y_labels, colors)):
            ax.bar(df[x_col], df[y_col], color=bar_color)
            ax.set_title(title, fontsize=10)
            ax.set_ylabel(y_label)
            ax.tick_params(axis="x", labelsize=7, rotation=45)
            ax.margins(y=0.15)
            
            # 添加数值标签
            for x, y in zip(df[x_col], df[y_col]):
                ax.text(x, y, f"{int(y)}", ha="center", va="bottom", fontsize=font_size)
        
        plt.tight_layout(rect=[0, 0, 1, 0.95])
        buf = BytesIO()
        plt.savefig(buf, format="png", dpi=dpi)
        plt.close()
        buf.seek(0)
        return buf
    
    @staticmethod
    def insert_charts(
            excel_path: str,
            df_summary_by_comm: pd.DataFrame,
            df_summary_by_process_pid: pd.DataFrame,
            df_proc_irq: pd.DataFrame
    ) -> None:
        wb = load_workbook(excel_path)

        # ---------- 图 1：Top10 running_s -> task_summary_by_comm!K2 ----------
        if not df_summary_by_comm.empty and "running_s" in df_summary_by_comm.columns:
            ws_comm = wb["task_summary_by_comm"]
            df_top10_comm = df_summary_by_comm.sort_values("running_s", ascending=False).head(10)
            
            buf = TraceAnalyzer.create_bar_chart(
                df=df_top10_comm, 
                x_col="comm", 
                y_col="running_s", 
                title="Top 10 Processes by Running Time",
                y_label="Running Time (s)",
                color="steelblue",
                dpi=120
            )
            
            img = Image(buf)
            ws_comm.add_image(img, "K2")

        # ---------- 图 2：每个 comm 的 PID 指标 -> task_summary_by_process_pid ----------
        if not df_summary_by_comm.empty and not df_summary_by_process_pid.empty:
            ws_pid = wb["task_summary_by_process_pid"]
            insert_row = 2
            chart_col_index = BASE_COL_INDEX
            df_pid = df_summary_by_process_pid.copy()

            # 动态构建可用的指标列表
            metrics = []
            if "preempt_cs_count" in df_pid.columns:
                metrics.append(("preempt_cs_count", "Preemptive CS Count"))
            if "preempt_time_s" in df_pid.columns:
                df_pid["preempt_time_ms"] = df_pid["preempt_time_s"] * 1000
                metrics.append(("preempt_time_ms", "Preempt Time (ms)"))
            if "migrate_count" in df_pid.columns:
                metrics.append(("migrate_count", "CPU Migration Count"))
            if "irq_total_time_s" in df_pid.columns:
                df_pid["irq_total_time_ms"] = df_pid["irq_total_time_s"] * 1000
                metrics.append(("irq_total_time_ms", "IRQ Total Time (ms)"))

            if metrics:
                df_top10_comm = df_summary_by_comm.sort_values("running_s", ascending=False).head(10)
                
                for comm in df_top10_comm["comm"]:
                    df_proc = df_pid[df_pid["comm"] == comm].copy()
                    if df_proc.empty:
                        continue
                    df_proc = df_proc.sort_values("pid")
                    rows = df_proc.to_dict("records")
                    chunks = [rows[i:i + MAX_PID_PER_CHART] for i in range(0, len(rows), MAX_PID_PER_CHART)]

                    for chunk_index, chunk in enumerate(chunks):
                        chunk_dfs = []
                        for metric, _ in metrics:
                            chunk_df = pd.DataFrame(chunk)[["pid", metric]]
                            chunk_df["pid"] = chunk_df["pid"].astype(str)
                            chunk_dfs.append(chunk_df)
                        
                        y_cols = [metric for metric, _ in metrics]
                        titles = [title for _, title in metrics]
                        y_labels = y_cols
                        
                        buf = TraceAnalyzer.create_subplots_chart(
                            df_list=chunk_dfs,
                            x_col="pid",
                            y_cols=y_cols,
                            titles=titles,
                            y_labels=y_labels,
                            main_title=f"{comm} - PID Metrics ({chunk_index + 1}/{len(chunks)})",
                            nrows=len(metrics),
                            ncols=1,
                            figsize=(8, max(4, 1.5 * len(metrics))),
                            color="royalblue"
                        )
                        
                        img = Image(buf)
                        col_letter = get_column_letter(chart_col_index)
                        ws_pid.add_image(img, f"{col_letter}{insert_row}")
                        chart_col_index += 12

                    insert_row += int(img.height / ROW_HEIGHT) + 2
                    chart_col_index = BASE_COL_INDEX

        # ---------- 图 3：IRQ Top10 -> proc_irq_detail!H2 ----------
        if not df_proc_irq.empty and "irq_time_s" in df_proc_irq.columns:
            ws_proc_irq = wb["proc_irq_detail"]
            df_proc_irq = df_proc_irq.copy()
            df_proc_irq["comm_pid"] = df_proc_irq["comm"] + "_" + df_proc_irq["pid"].astype(str)
            df_proc_irq["irq_time_ms"] = df_proc_irq["irq_time_s"] * 1000
            
            df_top10_irq_time = df_proc_irq.sort_values("irq_time_ms", ascending=False).head(10)
            df_top10_irq_count = df_proc_irq.sort_values("irq_count", ascending=False).head(10)
            
            irq_dfs = [df_top10_irq_time, df_top10_irq_count]
            y_cols = ["irq_time_ms", "irq_count"]
            titles = ["Top 10 IRQ Time by Process", "Top 10 IRQ Count by Process"]
            y_labels = ["IRQ Time (ms)", "IRQ Count"]
            
            buf = TraceAnalyzer.create_subplots_chart(
                df_list=irq_dfs,
                x_col="comm_pid",
                y_cols=y_cols,
                titles=titles,
                y_labels=y_labels,
                main_title="IRQ Statistics",
                nrows=2,
                ncols=1,
                figsize=(10, 6),
                color=["purple", "teal"]
            )
            
            img = Image(buf)
            ws_proc_irq.add_image(img, "H2")

        wb.save(excel_path)

    def advance_time(self, cpu: int, timestamp: Decimal) -> None:
        last_timestamp = self.last_ts[cpu]
        if not last_timestamp or timestamp < last_timestamp:
            self.last_ts[cpu] = timestamp
            return
        dt = timestamp - last_timestamp
        if self.irq_stack[cpu]:
            self.irq_stack[cpu][-1].acc_time += dt
        elif cpu in self.current_task:
            comm, pid = self.current_task[cpu]
            if pid != 0:
                self.stats[self.get_task_key(cpu, comm, pid)].running += dt
        self.last_ts[cpu] = timestamp

    def handle_soft_entry(self, cpu: int, timestamp: Decimal, vec: str, action: str) -> None:
        self.advance_time(cpu, timestamp)
        comm, pid = self.current_task.get(cpu, (f"swapper/{cpu}", 0))
        self.irq_stack[cpu].append(
            IRQFrame("soft", int(vec), action, timestamp, self.get_task_key(cpu, comm, pid))
        )

    def handle_soft_exit(self, cpu: int, timestamp: Decimal, *args) -> None:
        """处理软中断退出事件，忽略多余的参数"""
        self.advance_time(cpu, timestamp)
        if self.irq_stack[cpu]:
            frame = self.irq_stack[cpu].pop()
            if frame.task_key.pid != 0:
                s = self.stats[frame.task_key].irqs[("soft", frame.irq_id, frame.name)]
                s["count"] += 1
                s["time"] += frame.acc_time

    def handle_wakeup(self, timestamp: Decimal, comm: str, pid: str, cpu: str) -> None:
        pid = int(pid)
        cpu = int(cpu)
        if pid == 0:
            return
        self.pid_comm[pid] = comm
        if pid in self.offcpu and self.offcpu[pid].state == "sleeping":
            oc = self.offcpu[pid]
            dt = timestamp - oc.timestamp
            if dt > 0:
                comm_name = self.pid_comm.get(pid, comm)
                key = self.get_task_key(oc.cpu, comm_name, pid)
                self.stats[key].sleeping += dt
            self.offcpu[pid] = OffCPUState("runnable", timestamp, oc.cpu)
            return
        if pid in self.offcpu and self.offcpu[pid].state == "preempted":
            oc = self.offcpu[pid]
            dt = timestamp - oc.timestamp
            if dt > 0:
                comm_name = self.pid_comm.get(pid, comm)
                key = self.get_task_key(oc.cpu, comm_name, pid)
                self.stats[key].preempt_time_s += dt
            self.offcpu[pid] = OffCPUState("runnable", timestamp, oc.cpu)
            return
        self.offcpu[pid] = OffCPUState("runnable", timestamp, cpu)

    def handle_sched_switch(self, cpu: int, timestamp: Decimal, m) -> None:
        prev_comm, prev_pid, prev_state, next_comm, next_pid = m.groups()
        prev_pid, next_pid = int(prev_pid), int(next_pid)
        self.advance_time(cpu, timestamp)
        self.pid_comm[prev_pid] = prev_comm
        self.pid_comm[next_pid] = next_comm
        if prev_pid != 0:
            prev_state_cat = self.classify_prev_state(prev_state)
            self.offcpu[prev_pid] = OffCPUState(prev_state_cat, timestamp, cpu)
            key = self.get_task_key(cpu, prev_comm, prev_pid)
            self.stats[key].cs_count += 1
            if prev_state_cat == "preempted":
                self.stats[key].preempt_cs_count += 1
        if next_pid != 0 and next_pid in self.offcpu:
            oc = self.offcpu.pop(next_pid)
            dt = timestamp - oc.timestamp
            if dt > 0:
                comm = self.pid_comm.get(next_pid, next_comm)
                key = self.get_task_key(cpu, comm, next_pid)
                if oc.state == "sleeping":
                    self.stats[key].sleeping += dt
                elif oc.state == "preempted":
                    self.stats[key].preempt_time_s += dt
                elif oc.state == "runnable":
                    self.stats[key].runnable += dt
        self.current_task[cpu] = (next_comm, next_pid)

    def handle_migrate(self, m) -> None:
        comm, pid, orig_cpu, _ = m.groups()
        if int(pid) == 0:
            return
        self.stats[self.get_task_key(int(orig_cpu), comm, int(pid))].migrate_count += 1

    def handle_irq_entry(self, cpu: int, timestamp: Decimal, irq_id: str, name: str) -> None:
        self.advance_time(cpu, timestamp)
        comm, pid = self.current_task.get(cpu, (f"swapper/{cpu}", 0))
        self.irq_stack[cpu].append(
            IRQFrame("hard", int(irq_id), name, timestamp, self.get_task_key(cpu, comm, pid))
        )

    def handle_irq_exit(self, cpu: int, timestamp: Decimal, *args) -> None:
        """处理硬中断退出事件，忽略多余的参数"""
        self.advance_time(cpu, timestamp)
        if self.irq_stack[cpu]:
            frame = self.irq_stack[cpu].pop()
            if frame.task_key.pid != 0:
                s = self.stats[frame.task_key].irqs[("hard", frame.irq_id, frame.name)]
                s["count"] += 1
                s["time"] += frame.acc_time

    def parse_trace_file(self, path: str) -> Decimal:
        trace_end_ts = Decimal("0")
        getcontext().prec = 30
        
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                    
                # 提取CPU和时间戳
                m_cpu = CPU_RE.search(line)
                m_ts = TS_RE.search(line)
                if not m_cpu or not m_ts:
                    continue
                    
                cpu = int(m_cpu.group(1))
                timestamp = Decimal(m_ts.group(1))
                trace_end_ts = timestamp
                
                # 处理事件
                if "softirq_entry" in line:
                    m = SOFT_ENTRY_RE.search(line)
                    if m:
                        self.handle_soft_entry(cpu, timestamp, *m.groups())
                    continue
                if "softirq_exit" in line:
                    m = SOFT_EXIT_RE.search(line)
                    if m:
                        self.handle_soft_exit(cpu, timestamp, *m.groups())
                    continue
                if "sched_wakeup" in line:
                    m = WAKEUP_RE.search(line)
                    if m:
                        self.handle_wakeup(timestamp, *m.groups())
                    continue
                if "sched_waking" in line:
                    m = WAKING_RE.search(line)
                    if m:
                        self.handle_wakeup(timestamp, *m.groups())
                    continue
                if "sched_switch" in line:
                    m = SWITCH_RE.search(line)
                    if m:
                        self.handle_sched_switch(cpu, timestamp, m)
                    continue
                if "sched_migrate_task" in line:
                    m = MIGRATE_RE.search(line)
                    if m:
                        self.handle_migrate(m)
                    continue
                if "irq_handler_entry" in line:
                    m = IRQ_ENTRY_RE.search(line)
                    if m:
                        self.handle_irq_entry(cpu, timestamp, *m.groups())
                    continue
                if "irq_handler_exit" in line:
                    m = IRQ_EXIT_RE.search(line)
                    if m:
                        self.handle_irq_exit(cpu, timestamp, *m.groups())
                    continue
                self.advance_time(cpu, timestamp)
                    
        return trace_end_ts


    def build_dataframes(self) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        summary_rows = []
        irq_detail_rows = []
        proc_irq = defaultdict(lambda: {"count": 0, "time": Decimal("0")})
        
        # 遍历统计数据，构建原始数据行
        for tk, s in self.stats.items():
            if tk.pid == 0:
                continue
            irq_total_count = sum(v["count"] for v in s.irqs.values())
            irq_total_time = sum(v["time"] for v in s.irqs.values())
            
            # 添加任务摘要行
            summary_rows.append({
                "cpu": tk.cpu,
                "comm": tk.comm,
                "pid": tk.pid,
                "running_s": s.running,
                "sleeping_s": s.sleeping,
                "runnable_s": s.runnable,
                "cs_count": s.cs_count,
                "preempt_cs_count": s.preempt_cs_count,
                "preempt_time_s": s.preempt_time_s,
                "migrate_count": s.migrate_count,
                "irq_total_count": irq_total_count,
                "irq_total_time_s": irq_total_time
            })
            
            # 添加中断详情行
            for (irq_type, irq_id, name), v in s.irqs.items():
                irq_detail_rows.append({
                    "cpu": tk.cpu,
                    "comm": tk.comm,
                    "pid": tk.pid,
                    "irq_type": irq_type,
                    "irq_id": irq_id,
                    "irq_name": name,
                    "irq_count": v["count"],
                    "irq_time_s": v["time"]
                })
                
                # 累积进程中断统计
                key = (tk.comm, tk.pid, irq_type, irq_id, name)
                proc_irq[key]["count"] += v["count"]
                proc_irq[key]["time"] += v["time"]

        # 创建DataFrames
        df_summary = pd.DataFrame(summary_rows)
        df_irq_detail = pd.DataFrame(irq_detail_rows)
        df_proc_irq = pd.DataFrame([
            {
                "comm": comm,
                "pid": pid,
                "irq_type": irq_type,
                "irq_id": irq_id,
                "irq_name": name,
                "irq_count": v["count"],
                "irq_time_s": v["time"]
            }
            for (comm, pid, irq_type, irq_id, name), v in proc_irq.items()
        ])
        
        # 定义所有需要的列，确保即使数据为空也能创建正确的DataFrame
        summary_cols = ["cpu", "comm", "pid", "running_s", "sleeping_s", "runnable_s",
                        "cs_count", "preempt_cs_count", "preempt_time_s", "migrate_count",
                        "irq_total_count", "irq_total_time_s"]
        by_comm_cols = ["comm", "running_s", "sleeping_s", "runnable_s",
                        "cs_count", "preempt_cs_count", "preempt_time_s", "migrate_count",
                        "irq_total_count", "irq_total_time_s"]
        by_pid_cols = ["comm", "pid", "running_s", "sleeping_s", "runnable_s",
                        "cs_count", "preempt_cs_count", "preempt_time_s", "migrate_count",
                        "irq_total_count", "irq_total_time_s"]
        
        if df_summary.empty:
            df_summary = pd.DataFrame(columns=summary_cols)
            df_summary_by_comm = pd.DataFrame(columns=by_comm_cols)
            df_summary_by_process_pid = pd.DataFrame(columns=by_pid_cols)
        else:
            # 使用pandas groupby替代手动构建summary_by_comm字典
            df_summary_by_comm = df_summary.groupby("comm", as_index=False).agg({
                "running_s": "sum",
                "sleeping_s": "sum",
                "runnable_s": "sum",
                "cs_count": "sum",
                "preempt_cs_count": "sum",
                "preempt_time_s": "sum",
                "migrate_count": "sum",
                "irq_total_count": "sum",
                "irq_total_time_s": "sum"
            })
            
            df_summary_by_process_pid = df_summary.groupby(["comm", "pid"], as_index=False).agg({
                "running_s": "sum",
                "sleeping_s": "sum",
                "runnable_s": "sum",
                "cs_count": "sum",
                "preempt_cs_count": "sum",
                "preempt_time_s": "sum",
                "migrate_count": "sum",
                "irq_total_count": "sum",
                "irq_total_time_s": "sum"
            })
        
        # 优化Decimal到float的转换
        def convert_decimal_columns(df: pd.DataFrame) -> pd.DataFrame:
            """转换DataFrame中的Decimal类型列到float类型"""
            for col in df.columns:
                # 检查列中是否包含Decimal类型值
                if any(isinstance(x, Decimal) for x in df[col].dropna()):
                    df[col] = df[col].apply(lambda x: float(x) if isinstance(x, Decimal) else x)
            return df
        
        # 转换所有DataFrame中的Decimal列
        df_summary = convert_decimal_columns(df_summary)
        df_irq_detail = convert_decimal_columns(df_irq_detail)
        df_proc_irq = convert_decimal_columns(df_proc_irq)
        df_summary_by_comm = convert_decimal_columns(df_summary_by_comm)
        df_summary_by_process_pid = convert_decimal_columns(df_summary_by_process_pid)
        
        return df_summary, df_irq_detail, df_proc_irq, df_summary_by_comm, df_summary_by_process_pid


def run(input: str, output: str) -> None:
    if not os.path.exists(input):
        logging.error(f"The input file {input} does not exist")
        return
    tracer = TraceAnalyzer()
    trace_end_ts = tracer.parse_trace_file(input)
    try:
        for cpu in tracer.last_ts.keys():
            tracer.advance_time(cpu, trace_end_ts)
        (
            df_summary,
            df_irq_detail,
            df_proc_irq,
            df_summary_by_comm,
            df_summary_by_process_pid
        ) = tracer.build_dataframes()

        tracer.write_excel(
            df_summary,
            df_irq_detail,
            df_proc_irq,
            df_summary_by_comm,
            df_summary_by_process_pid,
            output
        )
        logging.info("The basic statistics table has been generated")

        tracer.insert_charts(
            output,
            df_summary_by_comm,
            df_summary_by_process_pid,
            df_proc_irq
        )
        logging.info("Graphical interface completed")
    except Exception as e:
        logging.error(e)


if __name__ == "__main__":
    input = "C:/Users/沈权/Downloads/cpu_trace_202512231401/cpu_trace_all"
    output = "C:/Users/沈权/Downloads/cpu_trace_202512231401/cpu_trace_all.xlsx"
    run(input, output)
